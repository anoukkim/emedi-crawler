"""
eMedi 허가정보 수집기 - 의료기기 허가정보(emedi.mfds.go.kr) 크롤링 · 병합
Dev: Yuri Kim | 2026-10

실행: python app.py
필요: pip install -r requirements.txt
"""

import sys, time, json, re, html, traceback, platform
from datetime import datetime
from pathlib import Path

import pandas as pd
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options

from recovery import PERMIT_SHEET, build_recovery_plan, detail_complete, load_existing_records
from models import expand_records

from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QLabel, QPushButton, QLineEdit, QTextEdit, QProgressBar,
    QFileDialog, QFrame, QStackedWidget, QMessageBox, QCheckBox,
    QScrollArea, QTableWidget, QTableWidgetItem, QHeaderView,
    QAbstractItemView, QSizePolicy, QStyledItemDelegate, QStyleOptionViewItem,
    QSizeGrip
)
from PyQt6.QtCore import Qt, QThread, pyqtSignal, QObject, QUrl, QEvent
from PyQt6.QtGui import QFont, QFontDatabase, QDesktopServices, QColor, QPainter

# ── Colors ────────────────────────────────────────────────
APP_NAME  = "eMedi 허가정보 수집기"
APP_DATE  = "2026-10"
PRIMARY   = "#0047ff"
SUCCESS   = "#00994f"
DANGER    = "#e5281b"
BG        = "#ffffff"
CARD      = "#ffffff"
BORDER    = "#0a0a0a"
ROW_BORDER = "#e3e3e3"
TEXT      = "#0a0a0a"
MUTED     = "#5c5c5c"
FAINT     = "#9a9a9a"
MONO      = '"IBM Plex Mono", "Consolas", "Menlo", monospace'
SANS      = '"IBM Plex Sans KR", "Malgun Gothic", "맑은 고딕", "Apple SD Gothic Neo", sans-serif'

BASE_URL   = "https://emedi.mfds.go.kr/search/data/MNU20237"

# ── 대기 설정 ─────────────────────────────────────────────
# 빠른 모드: 모달 내용이 나타나고 MODAL_SETTLE초 동안 변하지 않으면 바로 수집
MODAL_TIMEOUT = 20     # 모달 최대 대기 (초)
MODAL_SETTLE  = 1.0    # 내용이 이 시간 동안 그대로면 로딩 완료로 판단 (초)
PAGE_TIMEOUT  = 15     # 페이지 이동 최대 대기 (초)
POLL          = 0.25   # 확인 간격 (초)
POLITE_DELAY  = 0.3    # 항목 사이 쉬는 시간 (서버 부담 줄이기)
# 안전 모드 (기존 방식): 고정 대기
MODAL_WAIT = 10
RETRY_WAIT = 20
MAX_RETRY  = 2
EMPTY_WAIT = 10    # 모달 내용이 비어 있으면 추가로 기다리는 시간 (초)

LIST_COLS = [
    "순번","업체명","품목명","품목허가번호","품목등급","품목상태","취소취하일자",
    "품목분류번호","제품명","모델명","품목허가일자","품목일련번호","유효기간",
    "업허가번호","업종","업상태","업허가일자","요양급여코드",
    "추적관리대상여부","수출용에한함","위탁제조여부","유지관리용여부",
    "인체이식형여부","제조국가","제조자","제조의뢰국가명",
    "품목영문명","주소","UDI코드","포장수량",
]

MODAL_SEL = ".modal__frame.fit_modal_open"
RESIZE_MARGIN = 8   # 창 가장자리에서 크기 조절이 잡히는 폭 (px)

# 수집 엑셀 시트: 첫 시트는 모델명(UPN)별, 둘째 시트는 허가증 1건당 1행 (누락분 수집용)
SHEET_MODELS  = "모델명별"
SHEET_PERMITS = PERMIT_SHEET

CUR_PAGE_JS = """
var el = document.querySelector('.paging .num a[title="현재페이지"] b');
return el ? parseInt(el.innerText) : 0;
"""

FIRST_ROW_JS = """
var r = document.querySelector('table.scrollTable tbody tr');
return r ? r.innerText : '';
"""

CLOSE_MODAL_JS = """
var b = document.querySelector('#modalCloseBtn');
if (b) { try { b.click(); } catch(e) {} }
return !document.querySelector('""" + MODAL_SEL + """');
"""

PARSE_JS = """
var modal=document.querySelector('.modal__frame.fit_modal_open');
if(!modal)return{error:'no modal'};
var result={};
var sectionMap={'표준코드':'UDI','제품정보':'제품정보','수입업자':'수입업자','제조업자':'수입업자','기타정보':'기타정보','보험':'보험'};
function getPrefix(t){for(var k in sectionMap)if(t.indexOf(k)>=0)return sectionMap[k];return null;}
function cleanText(el){var c=el.cloneNode(true);c.querySelectorAll('span,img').forEach(function(s){s.remove();});return c.innerText.trim();}
modal.querySelectorAll('table').forEach(function(table){
    var caption=table.querySelector('caption');
    var title=caption?caption.innerText.trim():'';
    var prefix=getPrefix(title);if(!prefix)return;
    var colHeaders=[];
    table.querySelectorAll('tr').forEach(function(row){
        var ths=row.querySelectorAll('th'),tds=row.querySelectorAll('td');
        if(ths.length>0&&tds.length===0){colHeaders=Array.from(ths).map(function(th){return cleanText(th);});}
        else if(ths.length>0&&tds.length>0){for(var i=0;i<Math.min(ths.length,tds.length);i++){var k=cleanText(ths[i]),v=tds[i].innerText.trim();if(!k)continue;var key=prefix+'_'+k;
                if(key.indexOf('UDI_')===0){  // UDI 값엔 쉼표가 들어갈 수 있어 ; 로 구분 (빈 값도 자리 유지)
                    result[key]=(key in result)?result[key]+'; '+v:v;
                }else{
                    result[key]=result[key]?result[key]+', '+v:v;
                }}}
        else if(tds.length>0&&colHeaders.length>0){
            if(prefix==='UDI'&&tds.length>1){  // 모델명(UPN)별 행 분리용: UDI 표 한 행씩 보관
                var udiRow={};colHeaders.forEach(function(h,i){if(h)udiRow[h]=i<tds.length?tds[i].innerText.trim():'';});
                (result._UDI_rows=result._UDI_rows||[]).push(udiRow);
            }
            colHeaders.forEach(function(h,i){if(!h)return;var v=i<tds.length?tds[i].innerText.trim():'';
                var key=prefix+'_'+h;
                if(key.indexOf('UDI_')===0){  // UDI 값엔 쉼표가 들어갈 수 있어 ; 로 구분 (빈 값도 자리 유지)
                    result[key]=(key in result)?result[key]+'; '+v:v;
                }else{
                    result[key]=result[key]?result[key]+', '+v:v;
                }});}
    });
});
return result;
"""

def default_output_dir():
    """바탕화면 (OneDrive 바탕화면 포함), 없으면 홈 폴더"""
    home = Path.home()
    for p in (home / "Desktop", home / "OneDrive" / "Desktop", home / "OneDrive" / "바탕 화면"):
        if p.is_dir():
            return p
    return home


class Aborted(Exception):
    pass


# ═══════════════════════════════════════════════════════════
# Worker
# ═══════════════════════════════════════════════════════════
class CrawlWorker(QObject):
    log      = pyqtSignal(str, str)
    status   = pyqtSignal(str)
    progress = pyqtSignal(int)
    stats    = pyqtSignal(int, int, int, str)   # 전체, 완료, 실패, 남은 시간
    done     = pyqtSignal(bool, str)
    row_event = pyqtSignal(int, str, int, str)  # 순번, 허가번호, 항목 수, RUN/OK/FAIL
    checkpoint = pyqtSignal(int)

    def __init__(self, driver, output_dir, file_name, fast=True, existing_file=None):
        super().__init__()
        self.driver     = driver
        self.output_dir = Path(output_dir)
        self.file_name  = clean_file_name(file_name, default_crawl_name())
        self.fast       = fast
        self.existing_file = existing_file
        self._abort     = False
        self._page      = None
        self._orig_page = 1

    def abort(self):
        self._abort = True

    # ── 대기 헬퍼 ─────────────────────────────────────────
    def _sleep(self, sec):
        """중단 버튼에 바로 반응하는 sleep"""
        end = time.monotonic() + sec
        while True:
            if self._abort:
                raise Aborted()
            rem = end - time.monotonic()
            if rem <= 0:
                return
            time.sleep(min(POLL, rem))

    def _wait_until(self, fn, timeout):
        """fn()이 참이 될 때까지 대기. 성공 시 값, 시간 초과 시 None"""
        end = time.monotonic() + timeout
        while True:
            if self._abort:
                raise Aborted()
            try:
                v = fn()
                if v:
                    return v
            except Aborted:
                raise
            except Exception:
                pass
            if time.monotonic() >= end:
                return None
            time.sleep(POLL)

    def run(self):
        all_records, items, finished = [], [], False
        fatal_error = ""
        partial_path = unique_path(self.output_dir / f"{self.file_name}_partial.json")
        try:
            self.output_dir.mkdir(parents=True, exist_ok=True)
            existing = load_existing_records(self.existing_file) if self.existing_file else []
            if existing:
                self.log.emit(f"▶  기존 결과 {len(existing)}건 불러옴: {self.existing_file}", "section")
            mode = "빠른 모드" if self.fast else "안전 모드 (고정 대기)"
            self.log.emit(f"▶  목록 수집 시작  ·  {mode}", "section")
            items = self._collect_all_list()
            self.log.emit(f"  총 {len(items)}개 허가증 확인", "ok")
            if not items:
                self.done.emit(False, "목록이 비어 있습니다. 검색 결과가 보이는지 확인하세요.")
                return

            all_records, queue, overlap, old_only = build_recovery_plan(items, existing, LIST_COLS)
            self.log.emit(f"  기존 목록 일치 {overlap}건 · 상세 재수집 {len(queue)}건 · 이미 완료되어 건너뜀 {len(items)-len(queue)}건", "ok")
            if old_only:
                self.log.emit(f"  ⚠  현재 목록에 없는 기존 {old_only}건도 결과에 보존합니다.", "warn")

            self.log.emit(f"\n▶  상세 정보 수집 ({len(queue)}건)", "section")
            if not self.fast and queue:
                # 원래 코드와 동일: 1페이지로 돌아가서 시작
                self.driver.execute_script("getList(1)")
                self._sleep(1.5)
                self._orig_page = 1
            t0 = time.monotonic()
            self.stats.emit(len(all_records), sum(detail_complete(r) for r in all_records),
                            sum(r.get("오류", "") not in ("", "상세 정보 미수집") for r in all_records), "계산 중")

            for step, i in enumerate(queue, 1):
                idx, item = i + 1, items[i]
                if self._abort:
                    raise Aborted()
                permit_no = item.get("품목허가번호", f"row_{idx}")
                self.status.emit(f"[{idx}/{len(items)}]  {permit_no} 수집 중")
                self.row_event.emit(idx, permit_no, 0, "RUN")
                self.progress.emit(int((len(items)-len(queue)+step) / len(items) * 95))

                record = all_records[i]
                try:
                    if self.fast:
                        self._goto_page(item["_page"])
                        detail = self._fetch_modal(item)
                    else:
                        self._goto_page_original(item, idx)
                        detail = self._fetch_modal_original(item)
                    record.update(detail)
                    record.pop("오류", None)
                    self.log.emit(f"  [{idx:3d}/{len(items)}]  {permit_no}  ✓  ({len(detail)}개 항목)", "ok")
                    self.row_event.emit(idx, permit_no, len(detail), "OK")
                except Aborted:
                    raise
                except Exception as e:
                    self.log.emit(f"  [{idx:3d}/{len(items)}]  {permit_no}  →  {e}", "err")
                    self.row_event.emit(idx, permit_no, 0, "FAIL")
                    record["오류"] = str(e)
                finally:
                    if self.fast:
                        self._close_modal()
                    else:
                        self._close_modal_original()
                per = (time.monotonic() - t0) / step
                left = int(per * (len(queue) - step))
                eta = f"{left // 60}분" if left >= 60 else f"{left}초"
                self.stats.emit(len(all_records), sum(detail_complete(r) for r in all_records),
                                sum(r.get("오류", "") not in ("", "상세 정보 미수집") for r in all_records), eta)

                if self.existing_file or step % 25 == 0:
                    self._save_json(all_records, partial_path)
                    if step % 25 == 0 or self.existing_file and step == 1:
                        self.log.emit(f"  ★ 중간 저장: {partial_path}", "dim")
                        self.checkpoint.emit(sum(detail_complete(r) for r in all_records))

                if self.fast:
                    self._sleep(POLITE_DELAY)

            # ── 실패한 항목은 마지막에 한 번 더 시도 ──
            failed_idx = [i for i in queue if not detail_complete(all_records[i])]
            if failed_idx:
                self.log.emit(f"\n▶  실패한 {len(failed_idx)}건 다시 시도", "section")
                for n, i in enumerate(failed_idx, 1):
                    if self._abort:
                        raise Aborted()
                    item = items[i]
                    permit_no = item.get("품목허가번호", f"row_{i+1}")
                    self.status.emit(f"[다시 시도 {n}/{len(failed_idx)}]  {permit_no}")
                    self.row_event.emit(i + 1, permit_no, 0, "RUN")
                    try:
                        if self.fast:
                            self._goto_page(item["_page"])
                            detail = self._fetch_modal(item)
                        else:
                            self._goto_page_original(item, i + 1)
                            detail = self._fetch_modal_original(item)
                        record = all_records[i]
                        record.update(detail)
                        record.pop("오류", None)
                        self.log.emit(f"  {permit_no}  ✓  ({len(detail)}개 항목)", "ok")
                        self.row_event.emit(i + 1, permit_no, len(detail), "OK")
                    except Aborted:
                        raise
                    except Exception as e:
                        self.log.emit(f"  {permit_no}  →  {e}", "err")
                        self.row_event.emit(i + 1, permit_no, 0, "FAIL")
                    finally:
                        if self.fast:
                            self._close_modal()
                        else:
                            self._close_modal_original()
                    if self.existing_file:
                        self._save_json(all_records, partial_path)
                n_fail = sum(not detail_complete(r) for r in all_records)
                self.stats.emit(len(all_records), len(all_records) - n_fail, n_fail, "0초")

            finished = True

        except Aborted:
            self.log.emit(f"\n⏹  중단됨 (상세 수집 완료 {sum(detail_complete(r) for r in all_records)}건)", "warn")
        except Exception as e:
            fatal_error = str(e)
            self.log.emit(f"\n❌  오류: {e}", "err")
            self.log.emit(traceback.format_exc(), "dim")

        # ── 어떤 경우든 지금까지 모은 데이터는 저장 ──
        if not all_records:
            if not finished:
                self.done.emit(False, fatal_error or "저장할 데이터가 없습니다.")
            return
        try:
            self.log.emit("\n▶  엑셀 변환 중", "section")
            xlsx = self._to_excel(all_records)
            self._save_json(all_records, xlsx.with_suffix(".json"))
            if partial_path.exists():
                partial_path.unlink()
        except Exception as e:
            self.done.emit(False, f"저장 실패: {e}")
            return

        missing = sum(not detail_complete(r) for r in all_records)
        failed = sum(r.get("오류", "") not in ("", "상세 정보 미수집") for r in all_records)
        if missing:
            self.log.emit(f"  ⚠  상세 누락 {missing}건 (실패 {failed}건, 미시도 {missing-failed}건) → 누락분만 수집 가능", "warn")
        complete = sum(detail_complete(r) for r in all_records)
        self.stats.emit(len(all_records), complete, failed, "0초" if finished else "중단")
        if finished:
            self.progress.emit(100)
        else:
            self.progress.emit(int(complete / len(all_records) * 100))
        if finished:
            self.done.emit(True, str(xlsx))
        else:
            self.done.emit(False, f"중단됨 ({sum(detail_complete(r) for r in all_records)}건 상세 수집 완료, {len(all_records)}건 저장됨)\n{xlsx}")

    # ── 페이지 이동 ───────────────────────────────────────
    def _current_page(self):
        try:
            return self.driver.execute_script(CUR_PAGE_JS) or 0
        except Exception:
            return 0

    def _goto_page(self, page):
        cur = self._current_page()
        # cur == 0: 페이지 표시를 못 찾음 (1페이지뿐인 경우 등) → 마지막으로 이동한 페이지 기준
        if cur == page or (cur == 0 and self._page == page):
            return
        before = self.driver.execute_script(FIRST_ROW_JS)
        self.driver.execute_script("getList(arguments[0])", page)
        ok = self._wait_until(
            lambda: self._current_page() in (page, 0)
                    and self.driver.execute_script(FIRST_ROW_JS) not in ("", before),
            PAGE_TIMEOUT)
        if not ok:
            raise Exception(f"{page}페이지 이동 실패 (시간 초과)")
        self._page = page
        self._sleep(0.3)  # 행 렌더링 마무리

    # ── 상세 모달 ─────────────────────────────────────────
    def _fetch_modal(self, item):
        args = [item["_meddevseq"], item["_typename"], item["_index"],
                item["_mdentpsno"], item["_isfirst"], "", item["_seqgroup"]]

        for attempt in range(1, MAX_RETRY + 2):
            if attempt > 1:
                self.log.emit(f"    재시도 {attempt-1}회...", "warn")

            self._close_modal()
            # 인자를 문자열로 이어붙이지 않고 그대로 전달 (따옴표 등 특수문자 안전)
            self.driver.execute_script("itemDetail.apply(null, arguments)", *args)

            if self.fast:
                detail = self._wait_modal_settled(MODAL_TIMEOUT if attempt == 1 else MODAL_TIMEOUT * 2)
            else:
                self._sleep(RETRY_WAIT if attempt > 1 else MODAL_WAIT)
                detail = self.driver.execute_script(PARSE_JS) or {}
                if detail.get("error"):
                    detail = None
            if detail_complete(detail):
                return detail
            self.log.emit(f"    상세 화면 확인: {self._modal_diag()}", "dim")
            if attempt == 1:
                break   # 빠른 방식이 안 되면 바로 원래 방식으로
            self._close_modal()
            self._sleep(1)

        self.log.emit("    빠른 방식 실패 → 원래 방식(고정 대기)으로 재시도", "warn")
        self._close_modal_original()
        return self._fetch_modal_original(item)

    # ── 원래 방식 (처음 받은 코드와 동일한 동작) ──────────
    def _goto_page_original(self, item, idx):
        seq_num   = int(re.sub(r"\D", "", item.get("순번", "")) or str(idx))
        item_page = (seq_num - 1) // 10 + 1
        if item_page == self._orig_page:
            return
        self.driver.execute_script(f"getList({item_page})")
        for _ in range(5):          # 페이지 로딩 대기 (최대 5초)
            self._sleep(1)
            if self._current_page() == item_page:
                break
        self._orig_page = item_page

    def _fetch_modal_original(self, item):
        ms = item["_meddevseq"]
        tn = item["_typename"].replace("'", "\\'")
        mp = item["_mdentpsno"]
        fi = item["_isfirst"]
        sg = item["_seqgroup"]
        ix = item["_index"]

        for attempt in range(1, MAX_RETRY + 2):
            wait = RETRY_WAIT if attempt > 1 else MODAL_WAIT
            if attempt > 1:
                self.log.emit(f"    재시도 {attempt-1}회 ({wait}초 대기)...", "warn")
            self._close_modal_original()
            self._sleep(0.5)
            self.driver.execute_script(
                f"itemDetail('{ms}','{tn}','{ix}','{mp}','{fi}','','{sg}')")
            self._sleep(wait)
            detail = self.driver.execute_script(PARSE_JS) or {}
            if not detail.get("error") and not detail:
                # 모달은 열렸는데 내용이 아직 없음 (사이트가 느릴 때) → 최대 EMPTY_WAIT초 더 기다림
                detail = self._wait_until(
                    lambda: (self.driver.execute_script(PARSE_JS) or {}) or None, EMPTY_WAIT) or {}
            if detail_complete(detail):
                return detail
            self.log.emit(f"    상세 화면 확인: {self._modal_diag()}", "dim")
            self._close_modal_original()
            self._sleep(1)

        if detail.get("error"):
            raise Exception(f"상세 화면(모달)이 열리지 않음 ({MAX_RETRY}회 재시도)")
        raise Exception(f"상세 화면이 비어 있거나 내용이 불완전함 ({MAX_RETRY}회 재시도)")

    def _close_modal_original(self):
        try:
            self.driver.find_element(By.CSS_SELECTOR, "#modalCloseBtn").click()
            self._sleep(0.8)
        except Aborted:
            raise
        except Exception:
            pass

    def _modal_diag(self):
        """실패 원인 파악용: 지금 화면의 모달 상태를 한 줄로"""
        try:
            return self.driver.execute_script("""
                var all = document.querySelectorAll('.modal__frame');
                var open = document.querySelector('.modal__frame.fit_modal_open');
                var caps = open ? Array.from(open.querySelectorAll('caption')).map(function(c){return c.textContent.trim();}) : [];
                var nt = open ? open.querySelectorAll('table').length : 0;
                var head = open ? open.textContent.replace(/\\s+/g, ' ').trim().slice(0, 60) : '';
                return '모달 ' + all.length + '개, 열린 모달 ' + (open ? '있음' : '없음')
                     + (open ? ', 표 ' + nt + '개, 표 제목: ' + (caps.join(' / ') || '없음') + ', 내용: ' + head : '');
            """)
        except Exception as e:
            return f"확인 불가 ({str(e).splitlines()[0][:80]})"

    def _wait_modal_settled(self, timeout):
        """모달 내용이 나타나고 MODAL_SETTLE초 동안 바뀌지 않으면 반환"""
        end = time.monotonic() + timeout
        last, stable_since = None, None
        while time.monotonic() < end:
            if self._abort:
                raise Aborted()
            try:
                d = self.driver.execute_script(PARSE_JS) or {}
            except Exception:
                d = {}
            if d and not d.get("error"):
                if d == last:
                    if time.monotonic() - stable_since >= MODAL_SETTLE:
                        return d
                else:
                    last, stable_since = d, time.monotonic()
            time.sleep(POLL)
        return last  # 시간 초과: 마지막으로 본 내용이라도 반환 (없으면 None)

    def _close_modal(self):
        try:
            self._wait_until(lambda: self.driver.execute_script(CLOSE_MODAL_JS), 5)
        except Aborted:
            raise
        except Exception:
            pass

    # ── 목록 ──────────────────────────────────────────────
    def _collect_all_list(self):
        JS = """
        var rows=document.querySelectorAll('table.scrollTable tbody tr'),result=[];
        rows.forEach(function(row){
            var tds=row.querySelectorAll('td'),cells=Array.from(tds).map(function(td){return td.innerText.trim();});
            var link=row.querySelector('a[data-meddevseq]'),attrs={};
            if(link){attrs={meddevseq:link.getAttribute('data-meddevseq')||'',
            typename:link.getAttribute('data-typename')||'',
            mdentpsno:link.getAttribute('data-mdentpsno')||'',
            isfirst:link.getAttribute('data-isfirst')||'',
            seqgroup:link.getAttribute('data-seqgroup')||'',
            index:link.getAttribute('data-index')||''};}
            if(cells.length>0)result.push({cells:cells,attrs:attrs});
        });return result;
        """
        all_items = []
        last = self._get_last_page()
        self.log.emit(f"  총 {last or '?'}페이지", "dim")
        page = self._current_page() or 1
        self._page = page
        if page != 1:
            self._goto_page(1)
            page = 1
        while True:
            rows = self.driver.execute_script(JS)
            items = []
            for row in rows:
                cells, attrs = row.get("cells", []), row.get("attrs", {})
                if not cells or not any(cells): continue
                item = {col: (cells[i] if i < len(cells) else "") for i, col in enumerate(LIST_COLS)}
                for k in ["meddevseq","typename","mdentpsno","isfirst","seqgroup","index"]:
                    item[f"_{k}"] = attrs.get(k, "")
                item["_page"] = page  # 상세 수집 때 이 페이지로 이동
                items.append(item)
            if not items:
                if last:
                    raise Exception(f"{page}페이지 목록이 비어 있습니다. 네트워크/사이트 로딩 후 다시 시도하세요.")
                break
            all_items.extend(items)
            self.log.emit(f"  페이지 {page:3d}: {len(items):2d}건  (누적 {len(all_items):4d}건)", "dim")
            if last and page >= last: break
            if not last and len(items) < 10: break
            try:
                self._goto_page(page + 1)
            except Aborted:
                raise
            except Exception:
                if last:
                    raise
                break  # 마지막 페이지를 모르는 경우: 더 이상 넘어가지 않으면 종료
            page += 1
        return all_items

    def _get_last_page(self):
        try:
            els = self.driver.find_elements(By.XPATH, "//a[.//span[contains(text(),'맨 마지막')]]")
            if not els:
                return None
            m = re.search(r"getList\((\d+)\)", els[0].get_attribute("onclick") or "")
            return int(m.group(1)) if m else None
        except Exception:
            return None

    def _to_excel(self, records):
        clean = [{k:v for k,v in r.items() if not k.startswith("_")} for r in records]
        df = self._frame(clean)

        if "UDI_모델명" in df.columns:
            def uniq(val):
                if not val: return ""
                seen = []
                for m in str(val).split("; "):
                    m = m.strip()
                    if m and m not in seen: seen.append(m)
                return "|".join(seen)
            pos = df.columns.get_loc("UDI_모델명") + 1
            df.insert(pos, "모델명_고유", df["UDI_모델명"].apply(uniq))

        # 모델명(UPN)별로 한 행씩, 허가증 공통 항목은 모든 행에 복사
        df_models = self._frame(expand_records(records))

        path = unique_path(self.output_dir / f"{self.file_name}.xlsx")
        self.log.emit(f"  엑셀 저장 중... (모델명별 {len(df_models)}행 / 허가증별 {len(df)}행)", "dim")
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            df_models.to_excel(writer, sheet_name=SHEET_MODELS, index=False)
            df.to_excel(writer, sheet_name=SHEET_PERMITS, index=False)
        self.log.emit(f"  ✓  저장: {path}", "ok")
        return path

    def _frame(self, rows):
        df = pd.json_normalize(rows)

        # dict/list 타입 컬럼 → str 변환
        for col in df.columns:
            if df[col].apply(lambda x: isinstance(x, (dict, list))).any():
                df[col] = df[col].astype(str)

        df = df.fillna("")

        if len(df.columns) > 16384:
            self.log.emit(f"  ⚠  컬럼 {len(df.columns)}개 → 16384 초과분 잘라냄", "warn")
            df = df.iloc[:, :16384]
        return df

    def _save_json(self, records, path):
        tmp = Path(str(path) + ".tmp")
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(records, f, ensure_ascii=False, indent=2)
        tmp.replace(path)


def default_crawl_name():
    return f"eMedi_허가정보_{datetime.now():%Y%m%d}"


def default_merge_name():
    return f"eMedi_허가정보_최종_{datetime.now():%Y%m%d}"


def clean_file_name(name, fallback):
    """파일 이름에 쓸 수 없는 문자 제거, .xlsx 확장자는 떼기"""
    name = re.sub(r'[\\/:*?"<>|]', "_", (name or "").strip())
    if name.lower().endswith(".xlsx"):
        name = name[:-5]
    return name.strip(" .") or fallback


def unique_path(path):
    """같은 이름 파일이 있거나 엑셀에서 열려 있으면 _2, _3 ... 붙이기"""
    path = Path(path)
    if not path.exists():
        return path
    n = 2
    while True:
        p = path.with_name(f"{path.stem}_{n}{path.suffix}")
        if not p.exists():
            return p
        n += 1



def format_excel(ws, df, highlight_cols=None):
    """Excel 워크시트 포매팅"""
    from openpyxl.styles import (
        PatternFill, Font, Alignment, Border, Side, GradientFill
    )
    from openpyxl.utils import get_column_letter

    # 산출물의 기존 Excel 색상은 데스크톱 UI 리스타일과 무관하게 유지한다.
    HEADER      = "4F46E5"
    HEADER_HL   = "06B6D4"
    HIGHLIGHT   = "ECFEFF"   # 전자민원 컬럼 배경
    ROW_EVEN    = "F8FAFC"   # 짝수행
    ROW_ODD     = "FFFFFF"   # 홀수행
    BORDER_COL  = "E5E7EB"

    thin = Side(style="thin", color=BORDER_COL)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ── 헤더 스타일 ──
    for col_idx, col_name in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=col_idx)
        is_highlight = highlight_cols and col_name in highlight_cols

        cell.fill = PatternFill("solid", fgColor=HEADER_HL if is_highlight else HEADER)
        cell.font = Font(
            bold=True, color="FFFFFF",
            name="맑은 고딕", size=10
        )
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = border

    # ── 행 스타일 ──
    n_cols = len(df.columns)
    for row_idx in range(2, ws.max_row + 1):
        bg = ROW_EVEN if row_idx % 2 == 0 else ROW_ODD
        for col_idx in range(1, n_cols + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            col_name = df.columns[col_idx - 1]
            is_highlight = highlight_cols and col_name in highlight_cols

            cell.fill = PatternFill("solid", fgColor=HIGHLIGHT if is_highlight else bg)
            cell.font = Font(name="맑은 고딕", size=10)
            cell.alignment = Alignment(
                vertical="center", wrap_text=False,
                horizontal="center" if col_name in [
                    "순번","품목등급","품목상태","인체이식형여부","일회용",
                    "추적관리대상","수출용에한함","UDI_코드체계"
                ] else "left"
            )
            cell.border = border

    # ── 열 너비 자동 조정 ──
    for col_idx, col_name in enumerate(df.columns, 1):
        col_letter = get_column_letter(col_idx)
        header_len = len(str(col_name)) * 1.8
        # NaN/float 안전하게 str 변환 후 길이 계산
        sample = df.iloc[:50][col_name].fillna("").astype(str)
        lengths = sample.map(lambda x: len(x))
        max_data = lengths.max() if len(lengths) > 0 else 0
        if pd.isna(max_data):
            max_data = 0
        width = min(max(header_len, float(max_data) * 1.1, 8), 60)
        ws.column_dimensions[col_letter].width = width

    # ── 헤더 행 높이 ──
    ws.row_dimensions[1].height = 32

    # ── 틀 고정 (첫 행) ──
    ws.freeze_panes = "A2"

    # ── 자동 필터 ──
    ws.auto_filter.ref = ws.dimensions

class MergeWorker(QObject):
    log      = pyqtSignal(str, str)
    status   = pyqtSignal(str)
    progress = pyqtSignal(int)
    done     = pyqtSignal(bool, str)

    def __init__(self, crawl_file, permit_file, output_dir, file_name):
        super().__init__()
        self.crawl_file  = crawl_file
        self.permit_file = permit_file
        self.output_dir  = Path(output_dir)
        self.file_name   = clean_file_name(file_name, default_merge_name())

    def run(self):
        try:
            self.log.emit("▶  파일 로딩 중", "section")
            self.progress.emit(10)

            df = pd.read_excel(self.crawl_file, engine="openpyxl")
            self.log.emit(f"  크롤링 파일: {len(df)}행", "ok")
            if "품목허가번호" not in df.columns:
                raise Exception("크롤링 파일에 '품목허가번호' 열이 없습니다. 파일 1을 확인하세요.")
            self.progress.emit(30)

            # A~F 열을 한 번만 읽음. 허가번호 컬럼(두 번째)은 매칭 키, 나머지에 전자민원_ 접두사
            df2 = pd.read_excel(self.permit_file, engine="openpyxl", usecols="A:F")
            orig_cols = list(df2.columns)
            if len(orig_cols) < 2:
                raise Exception("허가(신고)사항 파일의 열이 2개 미만입니다. 파일 2를 확인하세요.")
            permit_no_col = orig_cols[1]
            merge_cols = {col: f"전자민원_{col}" for col in orig_cols if col != permit_no_col}
            df2 = df2.rename(columns=merge_cols)
            renamed_cols = list(merge_cols.values())
            self.log.emit(f"  허가(신고)사항 파일: {len(df2)}행  /  매칭 키: '{permit_no_col}'  /  추가 컬럼: {renamed_cols}", "ok")
            self.progress.emit(50)

            self.log.emit("▶  매칭 중", "section")
            df["_key"]  = df["품목허가번호"].astype(str).str.replace(r"\s", "", regex=True)
            df2["_key"] = df2[permit_no_col].astype(str).str.replace(r"\s", "", regex=True)
            df2 = df2.drop_duplicates(subset=["_key"], keep="first")

            df = df.merge(
                df2[["_key"] + renamed_cols],
                on="_key", how="left", indicator="_m"
            )
            matched = int((df["_m"] == "both").sum())
            df = df.drop(columns=["_key", "_m"])
            self.log.emit(f"  매칭 완료: {matched}/{len(df)}건", "ok")
            if matched == 0:
                self.log.emit("  ⚠  매칭된 건이 없습니다. 파일 2의 두 번째 열이 허가번호인지 확인하세요.", "warn")
            self.progress.emit(80)

            from openpyxl import load_workbook
            self.output_dir.mkdir(parents=True, exist_ok=True)
            out = unique_path(self.output_dir / f"{self.file_name}.xlsx")
            df_save = df.copy()
            for col in df_save.columns:
                if df_save[col].apply(lambda x: isinstance(x, (dict, list))).any():
                    df_save[col] = df_save[col].astype(str)
            df_save = df_save.fillna("")
            self.log.emit(f"  엑셀 저장 중... ({len(df_save)}행 × {len(df_save.columns)}열)", "dim")
            df_save.to_excel(out, index=False, engine="openpyxl")
            wb = load_workbook(out)
            format_excel(wb.active, df_save, highlight_cols=set(renamed_cols) if renamed_cols else None)
            wb.save(out)
            self.log.emit(f"  ✓  최종 저장: {out}", "ok")
            self.progress.emit(100)
            self.status.emit("완료")
            self.done.emit(True, str(out))

        except Exception as e:
            self.log.emit(f"❌  오류: {e}", "err")
            self.done.emit(False, str(e))


class ChromeOpener(QObject):
    """크롬 실행은 오래 걸릴 수 있으므로 별도 스레드에서. 결과는 시그널로 UI에 전달."""
    done = pyqtSignal(object, str)   # driver 또는 None, 오류 메시지

    def run(self):
        try:
            opts = Options()
            opts.add_argument("--start-maximized")
            opts.add_argument("--lang=ko-KR")
            # Selenium 4.11+ 내장 Selenium Manager가 크롬 버전에 맞는 드라이버를 자동으로 받음
            driver = webdriver.Chrome(options=opts)
            driver.implicitly_wait(0)   # 대기는 코드에서 직접 처리 (암묵적 대기는 느려짐)
            driver.get(BASE_URL)
            try:
                driver.maximize_window()
            except Exception:
                pass  # 창 관리자에서 최대화를 지원하지 않아도 연결 자체는 유지
            self.done.emit(driver, "")
        except Exception as e:
            self.done.emit(None, str(e).splitlines()[0] if str(e) else repr(e))


# ═══════════════════════════════════════════════════════════
# Main Window
# ═══════════════════════════════════════════════════════════
class DropLineEdit(QLineEdit):
    """파일을 끌어다 놓으면 경로가 입력되는 입력칸"""
    def __init__(self, *a, **kw):
        super().__init__(*a, **kw)
        self.setAcceptDrops(True)

    def dragEnterEvent(self, e):
        if e.mimeData().hasUrls():
            e.acceptProposedAction()

    def dropEvent(self, e):
        urls = e.mimeData().urls()
        if urls:
            self.setText(urls[0].toLocalFile())


class TitleBar(QFrame):
    """장식 없는 창도 제목 영역에서 이동/최대화할 수 있게 한다."""
    def mousePressEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            handle = self.window().windowHandle()
            if handle and handle.startSystemMove():
                event.accept()
                return
            self.window()._drag_offset = event.globalPosition().toPoint() - self.window().frameGeometry().topLeft()
        super().mousePressEvent(event)

    def mouseMoveEvent(self, event):
        window = self.window()
        if window._drag_offset is not None and event.buttons() & Qt.MouseButton.LeftButton and not window.isMaximized():
            window.move(event.globalPosition().toPoint() - window._drag_offset)
        super().mouseMoveEvent(event)

    def mouseReleaseEvent(self, event):
        self.window()._drag_offset = None
        super().mouseReleaseEvent(event)

    def mouseDoubleClickEvent(self, event):
        self.window()._toggle_maximize()
        event.accept()


class RuledRowDelegate(QStyledItemDelegate):
    """행 배경은 상태별로 유지하면서 얇은 가로선을 그린다."""
    def paint(self, painter, option, index):
        painter.fillRect(option.rect, index.data(Qt.ItemDataRole.BackgroundRole) or QColor("white"))
        inset = QStyleOptionViewItem(option)
        inset.rect = option.rect.adjusted(26 if index.column() != 3 else 0, 0,
                                          -26 if index.column() == 3 else 0, 0)
        super().paint(painter, inset, index)
        painter.save()
        painter.setPen(QColor(ROW_BORDER))
        painter.drawLine(option.rect.bottomLeft(), option.rect.bottomRight())
        painter.restore()


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(APP_NAME)
        self.setWindowFlag(Qt.WindowType.FramelessWindowHint)
        self.setMinimumSize(840, 620)
        self.setAttribute(Qt.WidgetAttribute.WA_Hover)   # 가장자리에서 크기 조절 커서 표시
        avail = QApplication.primaryScreen().availableGeometry() if QApplication.primaryScreen() else None
        w, h = 1264, 1180
        if avail:
            w, h = min(w, int(avail.width() * 0.96)), min(h, int(avail.height() * 0.94))
        self.resize(w, h)
        self.driver   = None
        self.worker   = None
        self._threads = []  # 완료 전까지 참조 유지
        self._last_crawl_file = None
        self._last_merge_file = None
        self._drag_offset = None
        self._active_log_row = None
        self._log_table_mode = "crawl"
        self._build_ui()
        QApplication.instance().installEventFilter(self)

    def _build_ui(self):
        central = QWidget()
        central.setObjectName("central")
        central.setStyleSheet(f"""
            #central {{ background:white; border:1px solid {TEXT}; }}
            QLabel {{ background:transparent; }}
            QToolTip {{ background:{TEXT}; color:white; border:1px solid {TEXT}; padding:6px; border-radius:0; }}
            QScrollBar:vertical {{ background:white; width:9px; border-left:1px solid {ROW_BORDER}; }}
            QScrollBar::handle:vertical {{ background:{MUTED}; min-height:25px; border-radius:0; }}
            QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical {{ height:0; }}
            QScrollBar::add-page:vertical, QScrollBar::sub-page:vertical {{ background:white; }}
        """)
        self.setCentralWidget(central)
        root = QVBoxLayout(central)
        root.setContentsMargins(0,0,0,0)
        root.setSpacing(0)

        # 검은 앱 타이틀 바 (기본 창 장식 대신 드래그/최소화/최대화/닫기 제공)
        title = TitleBar()
        title.setObjectName("titleStrip")
        title.setFixedHeight(36)
        title.setStyleSheet(f"#titleStrip {{ background:{TEXT}; color:white; }}")
        tl = QHBoxLayout(title)
        tl.setContentsMargins(28, 0, 0, 0)
        tl.setSpacing(0)
        cap = QLabel("EMEDI / 허가정보 수집기 — v1.0")
        cap.setStyleSheet(f"color:white; font-family:{MONO}; font-size:13px;")
        tl.addWidget(cap)
        tl.addStretch()
        for glyph, handler in (("−", self.showMinimized), ("□", self._toggle_maximize), ("×", self.close)):
            b = QPushButton(glyph)
            b.setFixedSize(42, 35)
            b.setStyleSheet("QPushButton { color:white; background:#0a0a0a; border:0; border-radius:0; font-size:17px; }"
                            "QPushButton:hover { background:#333333; }")
            b.clicked.connect(handler)
            tl.addWidget(b)
        root.addWidget(title)

        # 헤더의 탭은 실제 페이지와 동기화된다.
        nav = QFrame()
        nav.setFixedHeight(64)
        nav.setStyleSheet("background:white;")
        nl = QHBoxLayout(nav)
        nl.setContentsMargins(0, 0, 0, 0)
        nl.setSpacing(0)
        brand = QWidget()
        brand.setFixedWidth(276)
        brand.setStyleSheet(f"border-right:1px solid {TEXT}; border-bottom:0;")
        brand_row = QHBoxLayout(brand)
        brand_row.setContentsMargins(28, 0, 14, 0)
        brand_row.setSpacing(9)
        logo = QLabel("eMedi")
        logo.setStyleSheet(f"color:{TEXT}; font-family:{SANS}; font-weight:700; font-size:20px; border:0;")
        site = QLabel("emedi.mfds.go.kr")
        site.setStyleSheet(f"font-family:{MONO}; font-size:11px; color:{MUTED}; border:0;")
        brand_row.addWidget(logo)
        brand_row.addWidget(site)
        nl.addWidget(brand)
        self.tab_crawl = self._nav_button("01  크롤링", 0)
        self.tab_merge = self._nav_button("02  허가사항 병합", 1)
        nl.addWidget(self.tab_crawl)
        nl.addWidget(self.tab_merge)
        spacer = QWidget()
        spacer.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Preferred)
        nl.addWidget(spacer)
        self.lbl_chrome = QLabel()
        self.lbl_chrome.setFixedWidth(208)
        self.lbl_chrome.setAlignment(Qt.AlignmentFlag.AlignCenter)
        nl.addWidget(self.lbl_chrome)
        root.addWidget(nav)
        root.addWidget(self._rule(2))
        self._set_chrome_state(False)
        self.tabs = QStackedWidget()
        self.tabs.currentChanged.connect(self._tab_changed)

        stats = QFrame()
        stats.setFixedHeight(108)
        stats.setStyleSheet("background:white;")
        sl = QHBoxLayout(stats)
        sl.setContentsMargins(0, 0, 0, 0)
        sl.setSpacing(0)
        self.tile_total = self._tile("진행", "0/0", TEXT)
        self.tile_done = self._tile("완료", "0", TEXT)
        self.tile_fail = self._tile("실패 · 재시도 2회", "0", DANGER)
        self.tile_eta = self._tile("남은 시간", "—", TEXT)
        for tile in (self.tile_total, self.tile_done, self.tile_fail, self.tile_eta):
            sl.addWidget(tile, 1)
        self.tile_pending = self.tile_total.pending
        root.addWidget(stats)
        root.addWidget(self._rule(1))

        progress_row = QFrame()
        progress_row.setFixedHeight(52)
        progress_row.setStyleSheet("background:white;")
        pl = QHBoxLayout(progress_row)
        pl.setContentsMargins(28, 0, 28, 0)
        pl.setSpacing(18)
        self.lbl_status = QLabel("대기 중")
        self.lbl_status.setMinimumWidth(165)
        self.lbl_status.setStyleSheet(f"color:{TEXT}; font-family:{MONO}; font-size:13px; border:0;")
        self.progress = QProgressBar()
        self.progress.setFixedHeight(14)
        self.progress.setTextVisible(False)
        self.progress.setStyleSheet(f"QProgressBar {{ border:1px solid {TEXT}; background:white; border-radius:0; }}"
                                    f"QProgressBar::chunk {{ background:{PRIMARY}; border-radius:0; }}")
        self.lbl_pct = QLabel("0%")
        self.lbl_pct.setAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
        self.lbl_pct.setFixedWidth(45)
        self.lbl_pct.setStyleSheet(f"color:{TEXT}; font-family:{MONO}; font-size:13px; font-weight:600; border:0;")
        pl.addWidget(self.lbl_status)
        pl.addWidget(self.progress, 1)
        pl.addWidget(self.lbl_pct)
        root.addWidget(progress_row)
        root.addWidget(self._rule(2))

        self.tabs.addWidget(self._build_tab1())
        self.tabs.addWidget(self._build_tab2())
        self.tabs.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        root.addWidget(self.tabs)
        self._tab_changed(0)

        # 요약 테이블은 가용 높이를 전부 사용한다. 상세 진단 로그만 선택적으로 펼친다.
        self.log_table = QTableWidget(0, 4)
        self.log_table.setHorizontalHeaderLabels(["#", "허가번호", "항목", "상태"])
        for col in range(4):
            self.log_table.horizontalHeaderItem(col).setTextAlignment(
                Qt.AlignmentFlag.AlignVCenter |
                (Qt.AlignmentFlag.AlignRight if col >= 2 else Qt.AlignmentFlag.AlignLeft))
        self.log_table.setShowGrid(False)
        self.log_table.setAlternatingRowColors(False)
        self.log_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.log_table.setSelectionMode(QAbstractItemView.SelectionMode.NoSelection)
        self.log_table.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        self.log_table.setItemDelegate(RuledRowDelegate(self.log_table))
        self.log_table.verticalHeader().hide()
        self.log_table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Fixed)
        self.log_table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        self.log_table.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeMode.Fixed)
        self.log_table.horizontalHeader().setSectionResizeMode(3, QHeaderView.ResizeMode.Fixed)
        self.log_table.setColumnWidth(0, 96)
        self.log_table.setColumnWidth(2, 120)
        self.log_table.setColumnWidth(3, 120)
        self.log_table.horizontalHeader().setFixedHeight(34)
        self.log_table.setStyleSheet(f"""
            QTableWidget {{ background:white; border:0; color:{TEXT}; font-family:{MONO}; font-size:13px; }}
            QHeaderView::section {{ background:white; color:{MUTED}; border:0; border-bottom:1px solid {TEXT};
                padding-left:28px; font-family:{MONO}; font-size:11px; letter-spacing:1px; text-align:left; }}
        """)
        root.addWidget(self.log_table, 1)

        self.log_box = QTextEdit()
        self.log_box.setReadOnly(True)
        self.log_box.document().setMaximumBlockCount(250)
        self.log_box.setFixedHeight(170)
        self.log_box.setVisible(False)
        self.log_box.setStyleSheet(f"QTextEdit {{ background:white; color:{TEXT}; border:0; border-top:2px solid {TEXT};"
                                   f"font-family:{MONO}; font-size:12px; padding:8px 28px; border-radius:0; }}")
        root.addWidget(self.log_box)

        footer = QFrame()
        footer.setFixedHeight(40)
        footer.setStyleSheet(f"background:white; border-top:2px solid {TEXT};")
        fl = QHBoxLayout(footer)
        fl.setContentsMargins(28, 0, 28, 0)
        fl.setSpacing(20)
        fl.addWidget(self._link_btn("로그 지우기", self._clear_log))
        self.lbl_log_summary = QLabel("현재 작업만 표시됩니다")
        self.lbl_log_summary.setStyleSheet(f"color:{MUTED}; font-family:{MONO}; font-size:11px; border:0;")
        fl.addWidget(self.lbl_log_summary, 1)
        self.btn_log_toggle = self._link_btn("상세 로그 펼치기", self._toggle_log)
        self.btn_log_toggle.setCheckable(True)
        fl.addWidget(self.btn_log_toggle)
        credit = QLabel(f"DEV. YURI KIM · {APP_DATE}")
        credit.setStyleSheet(f"color:{MUTED}; font-family:{MONO}; font-size:11px; border:0;")
        fl.addWidget(credit)
        # 오른쪽 아래 모서리를 끌어서 창 크기 조절
        grip = QSizeGrip(footer)
        grip.setFixedSize(16, 16)
        grip.setStyleSheet("border:0; background:transparent;")
        fl.addWidget(grip, 0, Qt.AlignmentFlag.AlignBottom | Qt.AlignmentFlag.AlignRight)
        root.addWidget(footer)

    def _nav_button(self, label, index):
        b = QPushButton(label)
        b.setFixedWidth(184 if index else 132)
        b.setFixedHeight(62)
        b.clicked.connect(lambda: self.tabs.setCurrentIndex(index))
        return b

    def _tab_changed(self, index):
        self.tabs.setFixedHeight(165 if index == 0 else 252)
        for i, b in enumerate((self.tab_crawl, self.tab_merge)):
            color, bg = ("white", PRIMARY) if i == index else (TEXT, "white")
            b.setStyleSheet(f"QPushButton {{ background:{bg}; color:{color}; border:0; border-right:1px solid {TEXT};"
                            f"border-radius:0; font-family:{MONO}; font-size:12px; font-weight:600; }}"
                            f"QPushButton:hover {{ background:{PRIMARY if i == index else '#f2f2f2'}; }}")

    def _toggle_maximize(self):
        self.showNormal() if self.isMaximized() else self.showMaximized()

    def _rule(self, thickness):
        line = QFrame()
        line.setFixedHeight(thickness)
        line.setStyleSheet(f"background:{TEXT}; border:0;")
        return line

    def _resize_edge(self, pos):
        """창 가장자리(RESIZE_MARGIN px 안쪽)에 있으면 그 방향, 아니면 Qt.Edge(0)"""
        edge = Qt.Edge(0)
        if self.isMaximized():
            return edge
        if pos.x() < RESIZE_MARGIN: edge |= Qt.Edge.LeftEdge
        elif pos.x() >= self.width() - RESIZE_MARGIN: edge |= Qt.Edge.RightEdge
        if pos.y() < RESIZE_MARGIN: edge |= Qt.Edge.TopEdge
        elif pos.y() >= self.height() - RESIZE_MARGIN: edge |= Qt.Edge.BottomEdge
        return edge

    def eventFilter(self, obj, event):
        if obj is self and event.type() == QEvent.Type.HoverMove:
            edge = self._resize_edge(event.position().toPoint())
            E = Qt.Edge
            if edge in (E.LeftEdge | E.TopEdge, E.RightEdge | E.BottomEdge):
                self.setCursor(Qt.CursorShape.SizeFDiagCursor)
            elif edge in (E.RightEdge | E.TopEdge, E.LeftEdge | E.BottomEdge):
                self.setCursor(Qt.CursorShape.SizeBDiagCursor)
            elif edge in (E.LeftEdge, E.RightEdge):
                self.setCursor(Qt.CursorShape.SizeHorCursor)
            elif edge in (E.TopEdge, E.BottomEdge):
                self.setCursor(Qt.CursorShape.SizeVerCursor)
            else:
                self.unsetCursor()
        elif (event.type() == QEvent.Type.MouseButtonPress and
                event.button() == Qt.MouseButton.LeftButton and
                isinstance(obj, QWidget) and obj.window() is self):
            edge = self._resize_edge(self.mapFromGlobal(event.globalPosition().toPoint()))
            if edge and self.windowHandle() and self.windowHandle().startSystemResize(edge):
                return True
        return super().eventFilter(obj, event)

    # ── 탭 1: 크롤링 ──────────────────────────────────────
    def _build_tab1(self):
        w = QWidget()
        layout = QVBoxLayout(w)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)

        controls = QFrame()
        controls.setFixedHeight(56)
        controls.setStyleSheet("background:white;")
        row = QHBoxLayout(controls)
        row.setContentsMargins(0, 0, 0, 0)
        row.setSpacing(0)
        self.btn_open = self._btn("① 크롬 열기", CARD, self._open_site)
        self.btn_crawl = self._btn("② 크롤링 시작", PRIMARY, self._start_crawl)
        self.btn_crawl.setEnabled(False)
        self.btn_stop = self._btn("■  중단", DANGER, self._stop_crawl)
        self.btn_stop.setEnabled(False)
        for b, width in ((self.btn_open, 150), (self.btn_crawl, 165), (self.btn_stop, 125)):
            b.setFixedWidth(width)
            row.addWidget(b)

        self.chk_fast = QCheckBox("빠른 모드 (실험적)")
        self.chk_fast.setChecked(False)
        self.chk_fast.setToolTip("꺼짐 (기본): 원래 방식. 항목마다 상세 화면을 열고 10초 기다린 뒤 읽음\n"
                                 "켜짐: 상세 화면이 뜨면 바로 읽음. 실패하면 그 항목은 원래 방식으로 다시 시도")
        self.chk_fast.setStyleSheet(f"QCheckBox {{ color:{TEXT}; font-family:{SANS}; font-size:12px; spacing:8px; border:0; }}"
                                    f"QCheckBox::indicator {{ width:15px; height:15px; border:1px solid {TEXT}; border-radius:0; background:white; }}"
                                    f"QCheckBox::indicator:checked {{ background:{PRIMARY}; }}")
        fast = QWidget()
        fast.setStyleSheet(f"border-left:1px solid {TEXT}; border-bottom:0;")
        fast_row = QHBoxLayout(fast)
        fast_row.setContentsMargins(22, 0, 8, 0)
        fast_row.setSpacing(12)
        fast_row.addWidget(self.chk_fast)
        fast_hint = QLabel("기본 약 11초/항목")
        fast_hint.setStyleSheet(f"color:{MUTED}; font-family:{MONO}; font-size:11px; border:0;")
        fast_row.addWidget(fast_hint)
        fast_row.addStretch()
        row.addWidget(fast, 1)
        self.btn_crawl_open = self._ghost_btn("파일 열기 ↗", lambda: self._open_path(self._last_crawl_file))
        self.btn_crawl_dir = self._ghost_btn("폴더 열기 ↗",
            lambda: self._open_path(Path(self._last_crawl_file).parent if self._last_crawl_file else None))
        for b in (self.btn_crawl_open, self.btn_crawl_dir):
            b.setFixedWidth(118)
            row.addWidget(b)
        layout.addWidget(controls)
        layout.addWidget(self._rule(1))

        output = QFrame()
        output.setFixedHeight(56)
        output.setStyleSheet("background:white;")
        output_row = QHBoxLayout(output)
        output_row.setContentsMargins(0, 0, 0, 0)
        output_row.setSpacing(0)
        self.crawl_output, self.crawl_name = self._save_fields(output_row, default_crawl_name())
        layout.addWidget(output)
        layout.addWidget(self._rule(2))

        repair = QFrame()
        repair.setFixedHeight(48)
        repair.setStyleSheet("background:white;")
        repair_row = QHBoxLayout(repair)
        repair_row.setContentsMargins(28, 0, 0, 0)
        repair_row.setSpacing(0)
        repair_row.addWidget(self._field_label("기존 결과"))
        self.recovery_file = DropLineEdit()
        self.recovery_file.setPlaceholderText("기존 결과 .json / .xlsx (중간 저장 파일도 가능)")
        self._style_input(self.recovery_file)
        repair_row.addWidget(self.recovery_file, 1)
        repair_row.addWidget(self._small_btn("기존 결과 선택", self._pick_recovery))
        self.btn_repair = self._ghost_btn("누락분만 수집", lambda: self._start_crawl(repair=True))
        self.btn_repair.setEnabled(False)
        repair_row.addWidget(self.btn_repair)
        layout.addWidget(repair)
        layout.addWidget(self._rule(2))
        self._set_result_buttons("crawl", None)
        return w

    # ── 탭 2: 허가사항 병합 ───────────────────────────────
    def _build_tab2(self):
        w = QWidget()
        layout = QVBoxLayout(w)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)
        heading = QLabel("크롤링 결과에 전자민원 허가사항 붙이기  /  파일을 끌어다 놓거나 선택하세요")
        heading.setFixedHeight(40)
        heading.setStyleSheet(f"color:{TEXT}; padding-left:28px; font-family:{MONO}; font-size:12px;")
        layout.addWidget(heading)
        layout.addWidget(self._rule(1))
        row1 = self._merge_field("파일 1 · 크롤링 결과 엑셀")
        self.merge_crawl = DropLineEdit()
        self.merge_crawl.setPlaceholderText("여기에 파일을 끌어다 놓기")
        self._style_input(self.merge_crawl)
        row1.addWidget(self.merge_crawl)
        row1.addWidget(self._small_btn("선택", lambda: self._pick_file(self.merge_crawl)))
        layout.addLayout(row1)
        layout.addWidget(self._rule(1))

        row2 = self._merge_field("파일 2 · 품목 허가(신고)사항  (전자민원 → 나의민원 → 엑셀 다운로드)")
        self.merge_permit = DropLineEdit()
        self.merge_permit.setPlaceholderText("품목+허가(신고)사항.xlsx")
        self._style_input(self.merge_permit)
        row2.addWidget(self.merge_permit)
        row2.addWidget(self._small_btn("선택", lambda: self._pick_file(self.merge_permit)))
        layout.addLayout(row2)
        layout.addWidget(self._rule(1))

        output = QFrame()
        output.setFixedHeight(56)
        output.setStyleSheet("background:white;")
        ol = QHBoxLayout(output)
        ol.setContentsMargins(0, 0, 0, 0)
        ol.setSpacing(0)
        self.merge_output, self.merge_name = self._save_fields(ol, default_merge_name())
        layout.addWidget(output)
        layout.addWidget(self._rule(1))

        btn_row = QHBoxLayout()
        btn_row.setContentsMargins(0, 0, 0, 0)
        btn_row.setSpacing(0)
        self.btn_merge = self._btn("파일 병합 시작", PRIMARY, self._start_merge)
        self.btn_merge_open = self._ghost_btn("파일 열기 ↗", lambda: self._open_path(self._last_merge_file))
        self.btn_merge_dir  = self._ghost_btn("폴더 열기 ↗",
            lambda: self._open_path(Path(self._last_merge_file).parent if self._last_merge_file else None))
        btn_row.addWidget(self.btn_merge, 3)
        btn_row.addWidget(self.btn_merge_open, 1)
        btn_row.addWidget(self.btn_merge_dir, 1)
        layout.addLayout(btn_row)
        self._set_result_buttons("merge", None)
        return w

    def _merge_field(self, text):
        row = QHBoxLayout()
        row.setContentsMargins(28, 0, 0, 0)
        row.setSpacing(0)
        label = self._field_label(text)
        label.setMinimumWidth(460)
        row.addWidget(label)
        return row

    # ── UI 헬퍼 ───────────────────────────────────────────
    def _save_fields(self, layout, default_name):
        """저장 폴더 + 파일 이름 입력칸"""
        left = QWidget()
        left.setStyleSheet(f"border-right:1px solid {TEXT};")
        row_dir = QHBoxLayout(left)
        row_dir.setContentsMargins(28, 0, 0, 0)
        row_dir.setSpacing(0)
        row_dir.addWidget(self._field_label("저장 폴더"))
        out = DropLineEdit(str(default_output_dir()))
        self._style_input(out)
        row_dir.addWidget(out, 1)
        row_dir.addWidget(self._small_btn("변경", lambda: self._pick_dir(out)))

        right = QWidget()
        row_name = QHBoxLayout(right)
        row_name.setContentsMargins(16, 0, 28, 0)
        row_name.setSpacing(8)
        row_name.addWidget(self._field_label("파일명"))
        name = QLineEdit(default_name)
        name.setPlaceholderText(default_name)
        name.setToolTip("같은 이름의 파일이 있으면 _2, _3 … 이 붙어 저장됩니다.")
        self._style_input(name)
        ext = QLabel(".xlsx")
        ext.setStyleSheet(f"color:{MUTED}; font-family:{MONO}; font-size:11px;")
        row_name.addWidget(name, 1)
        row_name.addWidget(ext)
        layout.addWidget(left, 3)
        layout.addWidget(right, 2)
        return out, name

    def _field_label(self, text):
        l = QLabel(text)
        l.setStyleSheet(f"color:{MUTED}; font-family:{MONO}; font-size:11px; letter-spacing:1px; border:0;")
        l.setContentsMargins(0, 0, 16, 0)
        return l

    def _tile(self, label, value, color):
        t = QFrame()
        t.setStyleSheet(f"background:white; border-right:1px solid {TEXT};")
        v = QVBoxLayout(t)
        v.setContentsMargins(28, 12, 20, 6)
        v.setSpacing(2)
        cap = QLabel(label)
        cap.setStyleSheet(f"color:{MUTED}; font-family:{MONO}; font-size:11px; letter-spacing:1px; border:0;")
        val = QLabel(value)
        val.setStyleSheet(f"color:{color}; font-family:{MONO}; font-size:44px; font-weight:500; border:0;")
        v.addWidget(cap)
        v.addWidget(val)
        t.value = val
        if label == "진행":
            t.pending = QLabel("대기 0")
            t.pending.setStyleSheet(f"color:{MUTED}; font-family:{MONO}; font-size:10px; border:0;")
            v.addWidget(t.pending)
        return t

    def _btn(self, text, color, slot):
        b = QPushButton(text)
        b.setFixedHeight(56)
        b.setCursor(Qt.CursorShape.PointingHandCursor)
        fg = "white" if color != CARD else TEXT
        hover = self._dk(color) if color != CARD else "#f2f2f2"
        b.setStyleSheet(f"""
            QPushButton {{
                background:{color}; color:{fg};
                font-family:{SANS}; font-size:13px; font-weight:700;
                border-radius:0; border:0; border-right:1px solid {TEXT};
            }}
            QPushButton:hover {{ background:{hover}; }}
            QPushButton:disabled {{ background:{color}; color:{FAINT}; }}
        """)
        b.clicked.connect(slot)
        return b

    def _ghost_btn(self, text, slot):
        b = QPushButton(text)
        b.setFixedHeight(56)
        b.setCursor(Qt.CursorShape.PointingHandCursor)
        b.setStyleSheet(f"""
            QPushButton {{
                background:white; color:{TEXT}; font-family:{SANS};
                font-size:12px; font-weight:500;
                border-radius:0; border:0; border-left:1px solid {TEXT}; padding:0 14px;
            }}
            QPushButton:hover {{ background:#f2f2f2; }}
            QPushButton:disabled {{ color:{FAINT}; }}
        """)
        b.clicked.connect(slot)
        return b

    def _small_btn(self, text, slot):
        b = QPushButton(text)
        b.setFixedHeight(48)
        b.setCursor(Qt.CursorShape.PointingHandCursor)
        b.setStyleSheet(f"""
            QPushButton {{
                background:white; color:{PRIMARY}; font-family:{SANS};
                font-size:12px; font-weight:700;
                border-radius:0; border:0; border-left:1px solid {TEXT}; padding:0 16px;
            }}
            QPushButton:hover {{ background:#f2f2f2; }}
        """)
        b.clicked.connect(slot)
        return b

    def _link_btn(self, text, slot):
        b = QPushButton(text)
        b.setCursor(Qt.CursorShape.PointingHandCursor)
        b.setStyleSheet(f"""
            QPushButton {{ background:white; color:{TEXT}; border:0; border-radius:0;
                font-family:{MONO}; font-size:11px; }}
            QPushButton:hover {{ color:{PRIMARY}; }}
        """)
        b.clicked.connect(slot)
        return b

    def _style_input(self, edit):
        edit.setFixedHeight(44)
        edit.setStyleSheet(f"""
            QLineEdit {{
                background:white; border:0; border-radius:0; padding:0 12px;
                font-family:{MONO}; font-size:12px; color:{TEXT};
            }}
            QLineEdit:focus {{ border-bottom:2px solid {PRIMARY}; }}
        """)

    def _set_chrome_state(self, connected):
        dot, text = ("#00b35c", "CHROME · CONNECTED") if connected else (FAINT, "CHROME · DISCONNECTED")
        self.lbl_chrome.setText(f"<span style='color:{dot};'>■</span> &nbsp; {text}")
        self.lbl_chrome.setStyleSheet(f"""
            color:{MUTED}; background:white; border-left:1px solid {TEXT}; border-bottom:0;
            font-family:{MONO}; font-size:11px;
        """)

    def _set_result_buttons(self, which, path):
        """결과 파일이 생기면 '파일 열기' / '폴더 열기' 활성화"""
        if which == "crawl":
            self._last_crawl_file = path
            btns = (self.btn_crawl_open, self.btn_crawl_dir)
        else:
            self._last_merge_file = path
            btns = (self.btn_merge_open, self.btn_merge_dir)
        for b in btns:
            b.setEnabled(bool(path))
        btns[0].setToolTip(str(path) if path else "작업이 끝나면 결과 파일을 열 수 있습니다.")

    def _open_path(self, path):
        if path and Path(path).exists():
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(path)))
        elif path:
            self._log(f"파일을 찾을 수 없습니다: {path}", "err")

    def _set_stats(self, total, done, failed, eta):
        self.tile_total.value.setText(f"{done+failed:,}<span style='color:{FAINT};'>/{total:,}</span>")
        self.tile_done.value.setText(f"{done:,}")
        self.tile_fail.value.setText(f"{failed:03d}")
        self.tile_pending.setText(f"대기 {max(0, total-done-failed):,}")
        self.tile_eta.value.setText(eta)

    def _dk(self, h):
        r,g,b=int(h[1:3],16),int(h[3:5],16),int(h[5:7],16)
        return f"#{max(0,r-24):02x}{max(0,g-24):02x}{max(0,b-24):02x}"

    # ── 이벤트 ───────────────────────────────────────────
    def _pick_file(self, edit):
        p,_ = QFileDialog.getOpenFileName(self,"파일 선택",str(default_output_dir()),"Excel (*.xlsx *.xls)")
        if p: edit.setText(p)

    def _pick_recovery(self):
        start = self.recovery_file.text() or self.crawl_output.text() or str(default_output_dir())
        p,_ = QFileDialog.getOpenFileName(self, "기존 크롤링 결과 선택", start,
                                         "크롤링 결과 (*.json *.xlsx)")
        if p: self.recovery_file.setText(p)

    def _pick_dir(self, edit):
        p = QFileDialog.getExistingDirectory(self,"폴더 선택",edit.text() or str(Path.home()))
        if p: edit.setText(p)

    def _run_in_thread(self, worker):
        thread = QThread()
        worker.moveToThread(thread)
        thread.started.connect(worker.run)
        worker.done.connect(thread.quit)
        thread.finished.connect(lambda: self._forget_thread(thread))
        self._threads.append((thread, worker))
        thread.start()
        return thread

    def _forget_thread(self, thread):
        self._threads = [(t, w) for t, w in self._threads if t is not thread]
        thread.deleteLater()

    def _driver_alive(self):
        if not self.driver:
            return False
        try:
            self.driver.title  # 창이 닫혔으면 예외
            return True
        except Exception:
            return False

    def _open_site(self):
        if self._driver_alive():
            # 이미 열려 있으면 새 크롬을 띄우지 않고 검색 페이지로 이동
            self.driver.get(BASE_URL)
            try: self.driver.maximize_window()
            except Exception: pass
            self._log("이미 열린 크롬에서 검색 페이지로 이동했습니다.", "dim")
            return
        self.driver = None
        self.btn_open.setEnabled(False)
        self._log("크롬 여는 중... (처음 실행 시 드라이버 다운로드로 조금 걸릴 수 있어요)", "dim")
        opener = ChromeOpener()
        opener.done.connect(self._chrome_opened)
        self._run_in_thread(opener)

    def _chrome_opened(self, driver, err):
        self.btn_open.setEnabled(True)
        if driver is None:
            self._log(f"❌  크롬을 열지 못했습니다: {err}", "err")
            QMessageBox.warning(self, "크롬 실행 실패",
                "크롬을 열지 못했습니다.\n\n"
                "· Google Chrome이 설치되어 있는지 확인하세요.\n"
                "· 인터넷 연결을 확인하세요 (처음 실행 시 드라이버를 내려받습니다).\n\n"
                f"상세: {err}")
            return
        self.driver = driver
        self._set_chrome_state(True)
        self._log("✓  사이트 열림. 검색 후 크롤링 시작 버튼을 누르세요.", "ok")
        self.btn_crawl.setEnabled(True)
        self.btn_repair.setEnabled(True)

    def _start_crawl(self, repair=False):
        if not self._driver_alive():
            self.btn_crawl.setEnabled(False)
            self.btn_repair.setEnabled(False)
            self._set_chrome_state(False)
            self._log("크롬 창이 닫혀 있습니다. '크롬 열기'를 다시 누르고 검색하세요.", "err"); return
        existing_file = self.recovery_file.text().strip() if repair else None
        if repair and not Path(existing_file).is_file():
            QMessageBox.warning(self, "기존 결과 필요", "중간 저장 JSON 또는 기존 결과 Excel/JSON 파일을 선택하세요.")
            return
        self._clear_log()
        self._log_table_mode = "crawl"
        self.btn_crawl.setEnabled(False)
        self.btn_repair.setEnabled(False)
        self.btn_merge.setEnabled(False)
        self.btn_crawl.setText("누락분 수집 중..." if repair else "수집 중...")
        self.btn_stop.setEnabled(True)
        self.chk_fast.setEnabled(False)
        self._set_progress(0)
        self._set_stats(0, 0, 0, "—")
        self.tile_total.value.setText("목록…")

        worker = CrawlWorker(self.driver, self.crawl_output.text(), self.crawl_name.text(),
                             fast=self.chk_fast.isChecked(), existing_file=existing_file)
        worker.log.connect(self._log)
        worker.status.connect(self.lbl_status.setText)
        worker.progress.connect(self._set_progress)
        worker.stats.connect(self._set_stats)
        worker.row_event.connect(self._table_event)
        worker.checkpoint.connect(self._table_checkpoint)
        worker.done.connect(self._crawl_done)
        self.worker = worker
        self._run_in_thread(worker)

    def _stop_crawl(self):
        if self.worker:
            self.worker.abort()
            self._clear_log()
            self.btn_log_toggle.setChecked(False)
            self._log("중단 요청 중... 지금까지 모은 데이터를 저장합니다.", "warn")
            self.btn_stop.setEnabled(False)

    def _crawl_done(self, ok, msg):
        self.worker = None
        self.btn_crawl.setEnabled(self._driver_alive())
        self.btn_repair.setEnabled(self._driver_alive())
        self.btn_merge.setEnabled(True)
        self.btn_crawl.setText("② 크롤링 시작")
        self.btn_stop.setEnabled(False)
        self.chk_fast.setEnabled(True)
        saved = msg.splitlines()[-1]
        if saved.endswith(".xlsx"):
            self._set_result_buttons("crawl", saved)
            self.recovery_file.setText(str(Path(saved).with_suffix(".json")))
        if not self._driver_alive():
            self._set_chrome_state(False)
        if ok:
            self._log(f"\n🎉  크롤링 완료! → {msg}", "ok")
            self.lbl_status.setText("크롤링 완료")
            self.merge_crawl.setText(msg)   # 탭 2에 결과 파일 자동 입력
            self._ask_open("크롤링 완료", msg)
        else:
            self._log(f"\n⏹  {msg}", "warn")
            self.lbl_status.setText(msg.splitlines()[0])
            QMessageBox.information(self,"중단됨", msg)

    def _ask_open(self, title, path):
        r = QMessageBox.question(self, title, f"{title}!\n{path}\n\n파일을 지금 열까요?")
        if r == QMessageBox.StandardButton.Yes:
            self._open_path(path)

    def _start_merge(self):
        if self.worker:
            return  # 크롤링 중에는 같은 진행률/결과 표를 병합 작업과 공유하지 않는다.
        cf = self.merge_crawl.text().strip()
        pf = self.merge_permit.text().strip()
        od = self.merge_output.text().strip()
        if not cf or not pf:
            self._log("두 파일을 모두 선택하세요.", "err"); return
        for f in (cf, pf):
            if not Path(f).is_file():
                self._log(f"파일을 찾을 수 없습니다: {f}", "err"); return
        self._clear_log()
        self._log_table_mode = "merge"
        self.btn_merge.setEnabled(False)
        self.btn_crawl.setEnabled(False)
        self.btn_repair.setEnabled(False)
        self.btn_merge.setText("병합 중...")
        self._set_progress(0)

        worker = MergeWorker(cf, pf, od, self.merge_name.text())
        worker.log.connect(self._log)
        worker.status.connect(self.lbl_status.setText)
        worker.progress.connect(self._set_progress)
        worker.done.connect(self._merge_done)
        self._run_in_thread(worker)

    def _merge_done(self, ok, msg):
        self.btn_merge.setEnabled(True)
        self.btn_crawl.setEnabled(self._driver_alive())
        self.btn_repair.setEnabled(self._driver_alive())
        self.btn_merge.setText("파일 병합 시작")
        if ok:
            self._set_result_buttons("merge", msg)
            self._log(f"\n🎉  병합 완료! → {msg}", "ok")
            self._ask_open("병합 완료", msg)
        else:
            self._log(f"❌  {msg}", "err")
            QMessageBox.warning(self, "병합 실패", msg)

    def closeEvent(self, event):
        if self.worker:
            r = QMessageBox.question(self, "종료",
                "크롤링이 진행 중입니다. 중단하고 종료할까요?\n(지금까지 모은 데이터는 저장됩니다)")
            if r != QMessageBox.StandardButton.Yes:
                event.ignore(); return
            self.worker.abort()
        for t, _ in list(self._threads):
            t.quit(); t.wait(15000)
        if self.driver:
            try: self.driver.quit()   # 남아 있는 크롬/드라이버 프로세스 정리
            except Exception: pass
        event.accept()

    def _log(self, msg, level=""):
        colors = {"ok":SUCCESS,"err":DANGER,"warn":DANGER,
                  "section":PRIMARY,"dim":MUTED,"":TEXT}
        c = colors.get(level, MUTED)
        msg = msg.replace("🎉", "").replace("❌", "").replace("⏹", "").replace("⚠", "")
        msg = msg.replace("✓", "OK").replace("★", "-").replace("▶", "-")
        safe = html.escape(msg).replace("\n", "<br>")
        self.log_box.append(
            f'<span style="color:{c};">{safe}</span>')
        self.log_box.verticalScrollBar().setValue(
            self.log_box.verticalScrollBar().maximum())
        if level != "dim":
            brief = msg.strip().splitlines()[-1].strip() if msg.strip() else ""
            if brief:
                self.lbl_log_summary.setText(brief[:100] + ("…" if len(brief) > 100 else ""))
                self.lbl_log_summary.setToolTip(brief)
                if self._log_table_mode == "merge":
                    self._table_message(brief, level)
        try:
            with open(LOG_FILE, "a", encoding="utf-8") as f:
                f.write(f"{datetime.now():%Y-%m-%d %H:%M:%S}  {msg}\n")
        except Exception:
            pass

    def _clear_log(self):
        self.log_box.clear()
        self.log_table.setRowCount(0)
        self._active_log_row = None
        self.lbl_log_summary.setText("현재 작업만 표시됩니다")
        self.lbl_log_summary.setToolTip("")

    def _toggle_log(self, expanded):
        self.log_box.setVisible(expanded)
        self.btn_log_toggle.setText("상세 로그 접기" if expanded else "상세 로그 펼치기")

    def _table_event(self, index, permit, count, status):
        """화면에는 항목별 결과만, 상세 진단은 접힌 로그에 보관한다."""
        if status != "RUN" and self._active_log_row is not None:
            old = self._active_log_row
            if old < self.log_table.rowCount() and self.log_table.item(old, 1).text() == permit:
                row = old
            else:
                row = self._new_table_row()
        else:
            row = self._new_table_row()
        if status == "RUN":
            self._active_log_row = row
        else:
            self._active_log_row = None
        for col, value in enumerate((f"{index:03d}", permit, str(count) if count else "—",
                                      "수집 중…" if status == "RUN" else status)):
            item = QTableWidgetItem(value)
            item.setFlags(Qt.ItemFlag.ItemIsEnabled)
            item.setTextAlignment(Qt.AlignmentFlag.AlignVCenter |
                                  (Qt.AlignmentFlag.AlignRight if col >= 2 else Qt.AlignmentFlag.AlignLeft))
            color = "white" if status == "RUN" else (DANGER if status == "FAIL" and col == 3 else
                    SUCCESS if status == "OK" and col == 3 else FAINT if col == 0 else TEXT)
            item.setForeground(QColor(color))
            item.setBackground(QColor(TEXT if status == "RUN" else "white"))
            self.log_table.setItem(row, col, item)
        self.log_table.scrollToBottom()

    def _table_checkpoint(self, complete):
        row = self._new_table_row()
        for col, value in enumerate(("—", f"중간 저장 · {complete}건 기록됨", "", "")):
            item = QTableWidgetItem(value)
            item.setFlags(Qt.ItemFlag.ItemIsEnabled)
            item.setForeground(QColor(PRIMARY))
            item.setBackground(QColor("#eef2ff"))
            self.log_table.setItem(row, col, item)
        self.log_table.scrollToBottom()

    def _table_message(self, message, level):
        row = self._new_table_row()
        marker = "FAIL" if level == "err" else "OK" if level == "ok" else "—"
        for col, value in enumerate(("—", message, "", marker)):
            item = QTableWidgetItem(value)
            item.setFlags(Qt.ItemFlag.ItemIsEnabled)
            item.setForeground(QColor(DANGER if level == "err" else SUCCESS if level == "ok" and col == 3 else TEXT))
            self.log_table.setItem(row, col, item)
        self.log_table.setRowHeight(row, 35)
        self.log_table.scrollToBottom()

    def _new_table_row(self):
        if self.log_table.rowCount() >= 250:
            self.log_table.removeRow(0)
            if self._active_log_row is not None:
                self._active_log_row -= 1
        row = self.log_table.rowCount()
        self.log_table.insertRow(row)
        self.log_table.setRowHeight(row, 35)
        return row

    def _set_progress(self, v):
        self.progress.setValue(v)
        self.lbl_pct.setText(f"{v}%")


# ═══════════════════════════════════════════════════════════
LOG_FILE = Path.home() / "emedi_crawler.log"


def load_ui_fonts():
    """설치 여부에 의존하지 않도록 SIL OFL 글꼴을 앱에 등록한다."""
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent)) / "fonts"
    for name in (
        "IBMPlexSansKR-Regular.woff", "IBMPlexSansKR-Medium.woff", "IBMPlexSansKR-Bold.woff",
        "IBMPlexMono-Regular.woff", "IBMPlexMono-Medium.woff", "IBMPlexMono-SemiBold.woff",
    ):
        if (base / name).is_file():
            QFontDatabase.addApplicationFont(str(base / name))


def _excepthook(etype, value, tb):
    """앱으로 실행할 땐 콘솔이 없으므로, 예상 못한 오류를 창과 로그 파일로 보여줌"""
    text = "".join(traceback.format_exception(etype, value, tb))
    try:
        with open(LOG_FILE, "a", encoding="utf-8") as f:
            f.write(f"\n{datetime.now()}\n{text}\n")
    except Exception:
        pass
    sys.__excepthook__(etype, value, tb)
    if QApplication.instance():
        QMessageBox.critical(None, "오류", f"예상치 못한 오류가 발생했습니다.\n로그: {LOG_FILE}\n\n{value}")


if __name__ == "__main__":
    sys.excepthook = _excepthook
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    load_ui_fonts()
    app.setFont(QFont("IBM Plex Sans KR" if QFont("IBM Plex Sans KR").exactMatch() else
                      ("Apple SD Gothic Neo" if platform.system()=="Darwin" else "맑은 고딕"), 11))
    app.setApplicationName(APP_NAME)
    win = MainWindow()
    win.show()
    sys.exit(app.exec())
